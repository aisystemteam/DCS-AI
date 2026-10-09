from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
thin = Side(style="thin", color="999999")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
HDR_FILL = PatternFill("solid", fgColor="1F3864")
GRID_FILL = PatternFill("solid", fgColor="DDEBF7")
SUB_FILL = PatternFill("solid", fgColor="F2F2F2")
NOTE_FILL = PatternFill("solid", fgColor="FFF2CC")

def f(bold=False, color="000000", size=10):
    return Font(name=FONT, bold=bold, color=color, size=size)

# ---------------------------------------------------------------
# 패널 정의
#   grid : (row, col) 패널 상 물리적 위치 (좌상단 = 1,1)
#   positions : [(position명, 의미)]
# ---------------------------------------------------------------
AUDIO1 = {
    "panel": "AUDIO 1 CONTROL PANEL", "code": "AUD1", "tab": "AUDIO 1",
    "location": "Left Console",
    "source": "TO 1F-16CM/AM-1 BMS p.151-152 / DCS F-16C EA Guide p.65-66 대조 완료",
    "cols": 4, "rows": 2,
    "controls": [
        dict(no=1, grid=(1,1), cid="COMM1_PWR", name="COMM 1 (UHF) Power Knob", type="Rotary knob (power + volume)", push=None,
             positions=[("OFF", "UHF 무선기 전원 차단"),
                        ("INC (ON)", "CW 회전 시 전원 인가 및 UHF 오디오 볼륨 증가")],
             note="OFF 위치에서 CW 회전하면 전원 인가 (연속 가변)"),
        dict(no=2, grid=(1,2), cid="COMM2_PWR", name="COMM 2 (VHF) Power Knob", type="Rotary knob (power + volume)", push=None,
             positions=[("OFF", "VHF 무선기 전원 차단"),
                        ("INC (ON)", "CW 회전 시 전원 인가 및 VHF 오디오 볼륨 증가")],
             note="OFF 위치에서 CW 회전하면 전원 인가 (연속 가변)"),
        dict(no=3, grid=(1,3), cid="SECVOICE_VOL", name="SECURE VOICE Knob", type="Rotary knob (volume)", push=None,
             positions=[("NONE (OFF)", "최소 볼륨 (OFF 각인 없음)"), ("INC (ON)", "CW 회전 시 Secure Voice 오디오 볼륨 증가")],
             note="OFF 위치 없음 (볼륨 전용). DCS: No function"),
        dict(no=4, grid=(1,4), cid="MSL_VOL", name="MSL Tone Knob", type="Rotary knob (volume)", push=None,
             positions=[("NONE (OFF)", "최소 볼륨 (OFF 각인 없음)"), ("INC (ON)", "CW 회전 시 현재 선택된 AIM-9 미사일 톤 볼륨 증가")],
             note="OFF 위치 없음 (볼륨 전용)"),
        dict(no=8, grid=(2,1), cid="COMM1_MODE", name="COMM 1 (UHF) Mode Knob", type="3-position rotary knob (push)",
             push="PUSH: 수신 중단 후 선택 주파수로 Tone 신호 + HQ용 TOD 송신 (3개 위치 모두에서 가능)",
             positions=[("OFF", "Squelch 회로 비활성 — 약한 신호 수신 허용"),
                        ("SQL", "Squelch 회로 활성 — 정상 운용 시 배경 잡음 감소"),
                        ("GD", "UHF를 Guard 243.0 MHz로 동조, 전용 Guard 수신기 비활성 (C&I = BACKUP 시 GD 무효)")],
             note="PUSH 기능은 BMS TO 기준 — DCS 가이드 미언급"),
        dict(no=7, grid=(2,2), cid="COMM2_MODE", name="COMM 2 (VHF) Mode Knob", type="3-position rotary knob (push)",
             push="PUSH: VHF Tone 송신 (3개 위치 모두에서 가능)",
             positions=[("OFF", "Squelch 회로 비활성 — 약한 신호 수신 허용"),
                        ("SQL", "Squelch 회로 활성 — 배경 잡음 감소 (BMS: Not implemented yet)"),
                        ("GD", "VHF를 Guard 121.5 MHz로 동조")],
             note="PUSH 기능은 BMS TO 기준 — DCS 가이드 미언급. DCS: SQL 활성/비활성만 기술"),
        dict(no=6, grid=(2,3), cid="TF_VOL", name="TF Tone Knob", type="Rotary knob (volume)", push=None,
             positions=[("NONE (OFF)", "최소 볼륨 (OFF 각인 없음)"), ("INC (ON)", "CW 회전 시 TF(Terrain Following) 톤 볼륨 증가")],
             note="OFF 위치 없음 (볼륨 전용). DCS: No function"),
        dict(no=5, grid=(2,4), cid="THREAT_VOL", name="THREAT Tone Knob", type="Rotary knob (volume)", push=None,
             positions=[("NONE (OFF)", "최소 볼륨 (OFF 각인 없음)"), ("INC (ON)", "CW 회전 시 ALR-56M RWR 위협 경고음 볼륨 증가")],
             note="OFF 위치 없음 (볼륨 전용)"),
    ],
}

AUDIO2 = {
    "panel": "AUDIO 2 CONTROL PANEL", "code": "AUD2", "tab": "AUDIO 2",
    "location": "Left Console (AUDIO 1 하단)",
    "source": "TO 1F-16CM/AM-1 BMS p.151-153 / DCS F-16C EA Guide p.65-66 대조 완료",
    "cols": 4, "rows": 1,
    "controls": [
        dict(no=12, grid=(1,1), cid="INTERCOM_VOL", name="INTERCOM Knob", type="Rotary knob (volume)", push=None,
             positions=[("NONE (OFF)", "최소 볼륨 (OFF 각인 없음)"), ("INC (ON)", "CW 회전 시 인터컴(지상요원·급유 붐 오퍼레이터 통신) 볼륨 증가. 기어/저속 경고음·음성 메시지 볼륨에도 영향")],
             note="OFF 위치 없음 (볼륨 전용). DCS N/I"),
        dict(no=11, grid=(1,2), cid="TACAN_VOL", name="TACAN Knob", type="Rotary knob (volume)", push=None,
             positions=[("NONE (OFF)", "최소 볼륨 (OFF 각인 없음)"), ("INC (ON)", "CW 회전 시 TACAN 국 식별 신호 볼륨 증가")],
             note="TACAN 국 식별 모스 부호 청취용. BMS: TACAN 전원은 AVIONICS POWER 패널 MIDS LVT 노브로 인가"),
        dict(no=10, grid=(1,3), cid="ILS_PWR", name="ILS Power Knob", type="Rotary knob (power + volume)", push=None,
             positions=[("OFF", "ILS 수신기 전원 차단"),
                        ("INC (ON)", "CW 회전 시 전원 인가 및 ILS 로컬라이저 식별 모스 부호 볼륨 증가")],
             note="완전 CCW(OFF) 시 ILS 수신기 전원 차단"),
        dict(no=9, grid=(1,4), cid="HOTMIC_CIPHER", name="HOT MIC / CIPHER Switch", type="3-position toggle switch", push=None,
             positions=[("HOT MIC", "급유 붐 결합 시 지상요원/붐 오퍼레이터와 직접 통신 활성 (UHF/VHF 송신 시 우선 차단)"),
                        ("OFF", "HOT MIC·CIPHER 기능 비활성"),
                        ("CIPHER", "보안 음성 활성 시 UHF/VHF 비보안 신호 필터링")],
             note="DCS N/I"),
    ],
}

IFF = {
    "panel": "IFF CONTROL PANEL", "code": "IFF", "tab": "IFF",
    "location": "Left Console",
    "source": "TO 1F-16CM/AM-1 BMS p.148-150 / DCS F-16C EA Guide p.67-68 대조 완료",
    "cols": 5, "rows": 2,
    "controls": [
        dict(no=1, grid=(1,1), cid="MASTER", name="IFF MASTER Knob", type="5-position rotary knob (lift-to-turn at OFF/EMER)", push=None,
             positions=[("OFF", "트랜스폰더·IFF 질문 기능 전원 OFF (BMS: HOLD 미사용 시 Mode 4 제로화, OFF 전환 시 노브 들어올림)"),
                        ("STBY", "트랜스폰더 금지, IFF 질문기는 정상 동작"),
                        ("LOW", "트랜스폰더·질문 기능 운용 (NORM과 동일)"),
                        ("NORM", "트랜스폰더·질문 기능 운용"),
                        ("EMER", "노브를 들어올려 위치. Mode 1/2/3A 질문 인식 시마다 비상 펄스군 송신, Mode S 질문에 alert 응답")],
             note="APX-113 AIFF. C&I 노브 위치와 무관하게 동작. DCS N/I"),
        dict(no=2, grid=(1,2), cid="M4_CODE", name="IFF M-4 CODE Switch", type="3-position toggle switch (HOLD→A/B 스프링 복귀, ZERO 레버락)", push=None,
             positions=[("ZERO", "IFF MASTER가 OFF가 아닐 때 Mode 4 설정 제로화"),
                        ("A/B", "UFC 또는 (C&I=BACKUP 시) IFF MODE 4 REPLY 스위치로 코드 선택 허용 (중립/기본)"),
                        ("HOLD", "순간 위치. MASTER OFF 또는 전원 차단 전 HOLD로 두면 양쪽 코드 설정 유지")],
             note=""),
        dict(no=3, grid=(1,4), cid="CNI", name="C & I Knob", type="2-position rotary knob", push=None,
             positions=[("BACKUP", "UFC 고장 시 UHF 및 IFF 대체 운용 (UFC 정상 시에도 선택 가능)"),
                        ("UFC", "통신·항법·IFF를 UFC(Upfront Controls)로 정상 제어")],
             note="Communications and IFF 노브"),
        dict(no=7, grid=(2,1), cid="ENABLE", name="IFF ENABLE Switch", type="3-position toggle switch", push=None,
             positions=[("M3/MS", "BACKUP 시 Mode 3/A 및 S 활성"),
                        ("OFF", "BACKUP 시 Mode 1, 3/A, S 비활성"),
                        ("M1/M3", "BACKUP 시 Mode 1 및 3/A 활성")],
             note="C&I = BACKUP 시 모드 선택용"),
        dict(no=6, grid=(2,2), cid="MODE1_SEL", name="IFF MODE 1 Selector Levers", type="2 x thumbwheel lever (INC/DEC) + readout window", push=None,
             positions=[("2-digit code", "레버 증감으로 Mode 1 두 자리 코드 설정 (C&I=BACKUP 시)")],
             note="연속 숫자 선택기 — 개별 position 행 생략. DCS N/I"),
        dict(no=6, grid=(2,3), cid="MODE3_SEL", name="IFF MODE 3 Selector Levers", type="2 x thumbwheel lever (INC/DEC) + readout window", push=None,
             positions=[("2-digit code", "레버 증감으로 Mode 3/A 코드 상위 두 자리 설정, 하위 두 자리는 항상 00 (예: 77 → 7700)")],
             note="연속 숫자 선택기 — 개별 position 행 생략. DCS N/I"),
        dict(no=5, grid=(2,4), cid="M4_REPLY", name="IFF MODE 4 REPLY Switch", type="3-position toggle switch", push=None,
             positions=[("OUT", "Mode 4 운용 비활성"),
                        ("A", "Mode 4 활성, 프리셋 코드 A 선택"),
                        ("B", "Mode 4 활성, 프리셋 코드 B 선택")],
             note="C&I = BACKUP 시 사용. DCS N/I"),
        dict(no=4, grid=(2,5), cid="M4_MONITOR", name="IFF MODE 4 MONITOR Switch", type="2-position toggle switch", push=None,
             positions=[("AUDIO", "Mode 4 질문 감시음을 인터컴으로 제공 (미응답 시 0~1000 Hz 톤, 질문 수에 비례해 주파수 상승)"),
                        ("OUT", "오디오 모니터 비활성")],
             note="C&I = BACKUP 시 사용. DCS N/I"),
    ],
}

ANTSEL = {
    "panel": "ANTI ICE & ANT SEL PANEL", "code": "ANTSEL", "tab": "ANT SEL",
    "location": "Right Console",
    "source": "DCS F-16C EA Guide p.78 / TO 1F-16CM/AM-1 BMS p.157 대조 완료",
    "cols": 3, "rows": 1,
    "controls": [
        dict(no=1, grid=(1,1), cid="ANTI_ICE_ENG", name="ANTI-ICE ENGINE Switch", type="3-position toggle switch", push=None,
             positions=[("ON", "엔진 방빙 및 흡입구 스트럿 히터 수동 작동 (결빙 감지 시 INLET ICING caution은 여전히 점등)"),
                        ("AUTO", "흡입구 결빙 감지 시 방빙·히터 자동 작동 + INLET ICING caution 점등"),
                        ("OFF", "결빙 감지기·엔진 방빙·히터 비활성")],
             note="DCS 기준 ANT SEL과 같은 패널에 위치"),
        dict(no=2, grid=(1,2), cid="IFF_ANT", name="IFF ANT SEL Switch", type="3-position toggle switch", push=None,
             positions=[("UPPER", "상부 안테나로 질문 신호 수신 및 응답"),
                        ("NORM", "가장 강한 신호를 수신하는 안테나를 IFF가 자동 선택하여 응답"),
                        ("LOWER", "하부 안테나로 질문 신호 수신 및 응답")],
             note="IFF 질문 응답용 안테나 선택"),
        dict(no=3, grid=(1,3), cid="UHF_ANT", name="UHF ANT SEL Switch", type="3-position toggle switch", push=None,
             positions=[("UPPER", "상부 안테나로 송수신"),
                        ("NORM", "상·하부 안테나를 순환 사용하여 전방향 패턴 제공"),
                        ("LOWER", "하부 안테나로 송수신")],
             note=""),
    ],
}

UHFBU = {
    "panel": "UHF RADIO BACKUP CONTROL PANEL (HAVE QUICK II.2)", "code": "UHFBU", "tab": "UHF BACKUP",
    "location": "Left Console",
    "source": "TO 1F-16CM/AM-1 BMS p.159-160 (DCS EA Guide는 Radio Communications 장에 기술 — 본 범위 외, 미대조)",
    "cols": 5, "rows": 4,
    "controls": [
        dict(no=18, grid=(1,1), cid="ZERO", name="ZERO Switch", type="2-position toggle switch", push=None,
             positions=[("ZERO", "저장된 데이터 제로화 (TO 본문 설명 없음, 각인 기준)"),
                        ("(NORM)", "기본 위치")],
             note="접근 도어(1) 아래 위치"),
        dict(no=2, grid=(1,2), cid="FILL", name="FILL Connector", type="Connector (no position)", push=None,
             positions=[("—", "데이터 로드용 커넥터")], note="접근 도어 아래"),
        dict(no=3, grid=(1,3), cid="CHAN_DISP", name="CHAN Display", type="Indicator (readout)", push=None,
             positions=[("—", "선택된 프리셋 채널 번호 표시")], note=""),
        dict(no=5, grid=(1,4), cid="CHAN", name="CHAN Knob", type="20-position rotary knob", push=None,
             positions=[(str(i), f"프리셋 채널 {i} 선택 (Mode knob = PRESET 시 유효)") for i in range(1, 21)],
             note="TO: MWOD 시 1-19, single WOD 시 14 채널 사용 가능. WOD 저장 채널은 일반 프리셋으로 사용 불가. 채널 주파수는 도어의 카드에 수기 기재"),
        dict(no=17, grid=(2,1), cid="LOAD", name="LOAD Button", type="Push button (momentary)", push="PUSH: 현재 Manual 주파수를 CHAN 노브 채널의 프리셋으로 저장 (HQ II PRESET / HQ II.2 LOAD)",
             positions=[("PUSH", "프리셋 채널 주파수 저장")], note="접근 도어 아래. 절차: Function MAIN/BOTH → Mode PRESET → 수동 주파수 설정 → CHAN 선택 → 도어 개방 → LOAD"),
        dict(no=4, grid=(2,2), cid="FREQ_DISP", name="Frequency / Status Display", type="Indicator (readout)", push=None,
             positions=[("—", "주파수 또는 상태 표시")], note=""),
        dict(no=6, grid=(2,4), cid="STATUS", name="STATUS Button", type="Push button (momentary)", push="PUSH: 상태 표시",
             positions=[("PUSH", "Frequency/Status 디스플레이에 상태 표시")], note="TO 본문 설명 없음"),
        dict(no=16, grid=(3,1), cid="TEST_DISP", name="TEST DISPLAY Button", type="Push button (momentary)", push="PUSH: 디스플레이 테스트",
             positions=[("PUSH", "디스플레이 테스트")], note="TO 본문 설명 없음"),
        dict(no=15, grid=(3,2), cid="FREQ_100MHZ", name="A-3-2-T Knob (Manual Freq #1, 100 MHz)", type="4-position rotary knob", push=None,
             positions=[("T", "Test/Training 각인 (TO 설명 없음)"),
                        ("A", "HAVE QUICK Anti-jam 모드 (TO 설명 없음)"),
                        ("2", "2xx.xxx MHz"),
                        ("3", "3xx.xxx MHz")],
             note="수동 주파수 노브 5개 중 첫 번째 (225.000-399.975 MHz, 0.025 MHz 단위)"),
        dict(no=14, grid=(3,3), cid="FREQ_10MHZ", name="A-3-2 Knob (Manual Freq #2, 10 MHz)", type="Rotary knob (0-9 / A)", push=None,
             positions=[("0-9", "10 MHz 자리 숫자 선택")],
             note="TO 그림 각인 A-3-2 기준. 연속 숫자 선택기 — 개별 position 행 생략"),
        dict(no=7, grid=(3,4), cid="FREQ_1MHZ_KHZ", name="Manual Frequency Knobs #3-#5 (1 MHz / 100 kHz / 25 kHz)", type="3 x rotary knob (0-9, 0-9, 00/25/50/75)", push=None,
             positions=[("0-9", "1 MHz 자리"),
                        ("0-9", "100 kHz 자리"),
                        ("00/25/50/75", "25 kHz 자리")],
             note="5개 노브로 225.000-399.975 MHz를 0.025 MHz 단위로 수동 선택 (Mode knob = MANUAL 시 유효)"),
        dict(no=13, grid=(4,1), cid="FUNCTION", name="Function Knob", type="4-position rotary knob", push=None,
             positions=[("OFF", "전원 OFF"),
                        ("MAIN", "COMM 1 전원 스위치 ON 상태에서 선택 주파수로 UHF 운용"),
                        ("BOTH", "정상 운용 + Guard 주파수 동시 수신"),
                        ("ADF", "비동작 (Not operational)")],
             note="C&I = BACKUP 시에만 이 패널의 제어가 유효"),
        dict(no=12, grid=(4,2), cid="TONE", name="TONE Button / T-TONE Switch", type="Push button + toggle (HQ II.2: T-TONE switch [11], TONE button [12])", push="PUSH: 톤 송신",
             positions=[("TONE", "톤 송신 버튼 (TO 설명 없음)"),
                        ("T", "T-TONE 스위치 위치 (TO 설명 없음)")],
             note="HQ II 패널은 TONE 버튼만, HQ II.2 패널은 T-TONE 스위치(11) + TONE 버튼(12)"),
        dict(no=10, grid=(4,3), cid="VOL", name="VOL Knob", type="Rotary knob (INACTIVE)", push=None,
             positions=[("NONE (OFF)", "비동작"), ("INC (ON)", "비동작 — 볼륨은 AUDIO 1 패널 COMM 1 (UHF) Power Knob으로만 조절")],
             note="Inactive"),
        dict(no=9, grid=(4,4), cid="SQUELCH", name="SQUELCH Switch", type="2-position toggle switch", push=None,
             positions=[("OFF", "Squelch 비활성"),
                        ("ON", "Squelch 활성")],
             note="TO 본문 설명 없음, 각인 기준"),
        dict(no=8, grid=(4,5), cid="MODE", name="Mode Knob", type="3-position rotary knob", push=None,
             positions=[("MANUAL", "수동 주파수 노브 5개로 설정한 주파수 사용"),
                        ("PRESET", "CHAN 노브로 선택한 프리셋 주파수 사용"),
                        ("GUARD", "주 송수신기를 Guard 주파수로 자동 동조, Guard 수신기 비활성")],
             note=""),
    ],
}

def write_panel(wb, p):
    from openpyxl.comments import Comment
    ws = wb.create_sheet(p.get("tab") or p["panel"].replace(" CONTROL PANEL", ""))
    ws.sheet_view.showGridLines = False
    ws["A1"] = p["panel"]; ws["A1"].font = f(True, size=13)
    ws["A2"] = f"{p['location']}   |   {p['source']}"; ws["A2"].font = f(color="595959", size=9)
    grid_map = {ctl["grid"]: ctl for ctl in p["controls"]}
    ws.column_dimensions["A"].width = 3
    for c in range(1, p["cols"]+1):
        ws.column_dimensions[get_column_letter(2*c)].width = 22   # control
        ws.column_dimensions[get_column_letter(2*c+1)].width = 13 # positions
    top = 4
    for r in range(1, p["rows"]+1):
        h = max([len(grid_map[(r, c)]["positions"]) for c in range(1, p["cols"]+1) if (r, c) in grid_map] or [1])
        for c in range(1, p["cols"]+1):
            ctl = grid_map.get((r, c))
            cc, pc = 2*c, 2*c+1
            if ctl:
                cell = ws.cell(top, cc, ctl["name"] + "\n" + p["code"] + "." + ctl["cid"])
                cell.font = f(True); cell.fill = GRID_FILL
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
                tip = f"[{ctl['no']}] {ctl['type']}"
                if ctl.get("push"): tip += "\n" + ctl["push"]
                if ctl.get("note"): tip += "\n" + ctl["note"]
                cell.comment = Comment(tip, "TO")
                if h > 1:
                    ws.merge_cells(start_row=top, start_column=cc, end_row=top+h-1, end_column=cc)
                for k, (pos, desc) in enumerate(ctl["positions"]):
                    pcell = ws.cell(top+k, pc, pos); pcell.font = f(color="1F3864")
                    pcell.alignment = Alignment(vertical="center", wrap_text=True)
                    pcell.comment = Comment(desc, "TO")
            # borders: outline the control block and the position block
            for rr in range(top, top+h):
                for col in (cc, pc):
                    cell = ws.cell(rr, col)
                    cell.border = Border(left=thin, right=thin,
                                         top=thin if rr == top else None,
                                         bottom=thin if rr == top+h-1 else None)
                    if cell.font is None or not cell.value:
                        cell.font = f()
        top += h
    # notes block
    top += 1
    ws.cell(top, 2, "※ 셀 메모(코멘트)에 TO 기준 항목 번호·스위치 종류·각 Position 기능을 기재").font = f(color="595959", size=9)
    notes = [c for c in sorted(p["controls"], key=lambda x: x["grid"]) if c.get("note") or c.get("push")]
    for i, c in enumerate(notes, start=1):
        ws.cell(top+i, 2, f"[{c['no']}] {c['name']}").font = f(True, size=9)
        ws.cell(top+i, 4, " / ".join(x for x in (c.get("push"), c.get("note")) if x)).font = f(size=9)
    return ws


import re as _re
def pos_code(label):
    c = _re.sub(r"[^A-Za-z0-9]+", "_", label.upper()).strip("_")
    return c or "NA"

def flat_rows(panels):
    P, C, X = [], [], []
    for p in panels:
        P.append([p["code"], p["seq"], p["zone"], ZONE_RANK[p["zone"]], p["zone_seq"], p["panel"], p["location"], p["rows"], p["cols"], p["source"], p.get("note") or ""])
        for ctl in sorted(p["controls"], key=lambda x: x["grid"]):
            cid = f"{p['code']}.{ctl['cid']}"
            r, c = ctl["grid"]
            C.append([cid, p["code"], p["zone"], ZONE_RANK[p["zone"]], ctl["cid"], ctl["no"], ctl["name"], ctl["type"], r, c, f"R{r}C{c}",
                      len(ctl["positions"]), ctl.get("push") or "", ctl.get("note") or ""])
            seen = {}
            for i, (pos, desc) in enumerate(ctl["positions"], start=1):
                pc = pos_code(pos)
                if pc in seen:
                    seen[pc] += 1; pc = f"{pc}_{seen[pc]}"
                else:
                    seen[pc] = 1
                X.append([f"{cid}.{pc}", cid, p["code"], p["zone"], ZONE_RANK[p["zone"]], ctl["cid"], pc, i, pos, desc])
    return P, C, X

def write_table(wb, title, headers, rows, widths):
    ws = wb.create_sheet(title)
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(1, i, h); cell.font = f(True, "FFFFFF"); cell.fill = HDR_FILL
        cell.alignment = Alignment(vertical="center"); cell.border = box
        ws.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(rows, start=2):
        for i, v in enumerate(row, start=1):
            cell = ws.cell(r, i, v); cell.font = f(bold=(i == 1)); cell.border = box
            cell.alignment = Alignment(vertical="top", wrap_text=(i >= len(row) - 1))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows)+1}"
    return ws

from dcs_panels import DCS_PANELS
from icp import ICP, DED_PAGES, DED_TRANSITIONS, DED_FIELDS, ICP_ACTIONS
from mfd import MFDL, MFDR, MFD_FORMATS, MFD_OSB_MAP, MFD_TRANSITIONS, MFD_ACTIONS
from cmds import CMDS, CMDS_DED_PAGES, CMDS_DED_TRANSITIONS, CMDS_DED_FIELDS, CMDS_ACTIONS
from defensive import RWRAZ, RWRPRIME, RWRAUX, ECM, CFBTN, DEF_ACTIONS
from hotas import SSC, THROTTLE, HOTAS_FUNCTIONS, HOTAS_ACTIONS
from uhfbu import UHFBU as UHFBU_DCS, UHFVHF_DED_FIELDS, UHFBU_ACTIONS, ICP_RADIO_ACTIONS
UHFBU = UHFBU_DCS   # BMS 정의를 DCS 정의로 교체 (BMS HQ II.2 정의는 build_audio.py 상단에 보존)
PANELS = [IFF, AUDIO1, AUDIO2, ANTSEL, UHFBU] + DCS_PANELS + [ICP, MFDL, MFDR, CMDS, RWRAZ, RWRPRIME, RWRAUX, ECM, CFBTN, SSC, THROTTLE]
# CMDS DED 페이지: icp.py의 CMDS 스텁을 상세 페이지로 교체
DED_PAGES = [p for p in DED_PAGES if p[0] != "CMDS"]
_i = next(i for i, p in enumerate(DED_PAGES) if p[0] == "MODE")
DED_PAGES[_i:_i] = CMDS_DED_PAGES
DED_TRANSITIONS = [t if not (t[1] == "LIST" and t[2] == "ICP.KEY_7") else [t[0], "LIST", "ICP.KEY_7", "CMDS_BINGO", ""] for t in DED_TRANSITIONS]
DED_TRANSITIONS += [[f"TR{len(DED_TRANSITIONS)+i+1:03d}"] + r for i, r in enumerate(CMDS_DED_TRANSITIONS)]
DED_FIELDS = [f for f in DED_FIELDS if f[1] not in ("UHF", "VHF")] + UHFVHF_DED_FIELDS + CMDS_DED_FIELDS
DED_PAGES = [[p[0], p[1], p[2], p[3], p[4], p[5], "Y", "253-256", "Radio Communications 장으로 필드 보완 완료"] if p[0] in ("UHF", "VHF") else p for p in DED_PAGES]
# 패널 조작 매크로(비-ICP/MFD)는 PANEL_ACTIONS로 분리
PANEL_ACTIONS = HOTAS_ACTIONS + [a for a in CMDS_ACTIONS if not a[0].startswith("ICP.")] + DEF_ACTIONS + UHFBU_ACTIONS
# HOTAS ID 연결: MFD DMS 전이
MFD_TRANSITIONS = [[t[0], t[1], "SSC.DMS.LEFT / SSC.DMS.RIGHT (short)", t[3], t[4]] if t[2].startswith("HOTAS DMS") else t for t in MFD_TRANSITIONS]
ICP_ACTIONS = [a for a in ICP_ACTIONS if a[0] not in ("ICP.TUNE_UHF_MANUAL", "ICP.TUNE_UHF_PRESET")] + ICP_RADIO_ACTIONS + [a for a in CMDS_ACTIONS if a[0].startswith("ICP.")]

# ── 인체공학 접근 우선순위 (호야 지정): HOTAS > 전방(계기판/HUD/ICP/MFD) > 좌측 Aux > 좌측 콘솔 > 우측 Aux > 우측 콘솔
ZONES = [  # zone_id, access_rank, zone_name, 설명
    ("HOTAS", 1, "Hands-On Controls (Stick/Throttle)", "손을 떼지 않고 조작 — 전투 기동 중 사용"),
    ("FRONT", 2, "Front / Instrument Panel (HUD, ICP, MFD, 계기)", "전방 시선 유지한 채 조작 — 비행 중 상시"),
    ("LAUX",  3, "Left Auxiliary Console", "기어·훅·CMDS 등 이착륙/방어 — 왼손 근거리"),
    ("LCONS", 4, "Left Console", "스로틀 손 — 통신·연료·엔진·비행제어·조명"),
    ("RAUX",  5, "Right Auxiliary Console", "주의등·계기 — 주로 읽기"),
    ("RCONS", 6, "Right Console", "시동/전원/항전 설정 — 주로 지상·시동 단계에서 1회 세팅"),
]
ZONE_RANK = {z[0]: z[1] for z in ZONES}
def zone_of(loc):
    if loc.startswith("Left Auxiliary"): return "LAUX"
    if loc.startswith("Right Auxiliary"): return "RAUX"
    if loc.startswith("Left Console"): return "LCONS"
    if loc.startswith("Right Console"): return "RCONS"
    if loc.startswith("Instrument Panel"): return "FRONT"
    if loc.startswith("HOTAS"): return "HOTAS"
    raise ValueError(loc)
for p in PANELS:
    p["zone"] = zone_of(p["location"])
# 구역 내 물리 순서 (DCS EA Guide 개요 사진 기준, 전방→후방 / 상→하 / 좌→우)
ZONE_ORDER = {
    "HOTAS": ["SSC", "THROTTLE"],
    "FRONT": ["LEYEBROW", "RWRAZ", "RWRPRIME", "AOAIDX", "ARNWS", "REYEBROW", "ICP", "MFDL", "MFDR", "MISC", "INSTR", "FUELQTY"],          # p.44 (HUD/MFD 추후 삽입)
    "LAUX":  ["LG", "RWRAUX", "ALTGEAR", "CMDS", "HMCS"],                                                                 # p.55 (THREAT WARNING AUX/CMDS 추후)
    "LCONS": ["CANOPY", "CFBTN", "FLTCTL", "TEST", "IFF", "TRIM", "EXTLT", "AVTR", "ECM", "AUD1", "AUD2", "ELEC", "FUEL", "EPU", "ENGJET", "UHFBU", "MPO"],  # p.62
    "RAUX":  ["RAUXINST", "CAUTION"],                                                                   # p.58
    "RCONS": ["SNSRPWR", "RCONSMISC", "INTLT", "AIRCOND", "PLAINCIPHER", "OXYGEN", "ANTSEL", "KY58", "AVPWR"],  # p.75 (HUD 제어판 추후 맨 앞)
}
def _order(p):
    lst = ZONE_ORDER[p["zone"]]
    assert p["code"] in lst, f"ZONE_ORDER에 없음: {p['code']} ({p['zone']})"
    return (ZONE_RANK[p["zone"]], lst.index(p["code"]))
PANELS.sort(key=_order)
for i, p in enumerate(PANELS, start=1):
    p["seq"] = i
    p["zone_seq"] = ZONE_ORDER[p["zone"]].index(p["code"]) + 1
P, C, X = flat_rows(PANELS)

wb = Workbook()
ws0 = wb.active; ws0.title = "README"; ws0.sheet_view.showGridLines = False
lines = [
    ("F-16C (Falcon BMS) 조종석 패널 DB — Project Duality", True),
    ("", False),
    ("1. 구조 (LLM 조회용 정규화 테이블 + 사람 검수용 격자 탭)", True),
    ("  ZONES     : 조종석 구역 6개와 인체공학 접근 우선순위(access_rank 1=가장 쉬움). HOTAS > FRONT > LAUX > LCONS > RAUX > RCONS", False),
    ("              모든 표와 GRID 탭은 이 순위로 정렬. PANELS.seq = 전체 순번, zone_seq = 구역 내 물리 순서(전방→후방, 가이드 개요 사진 기준)", False),
    ("  PANELS    : 패널 1행 = 패널 1개. panel_id가 기본키", False),
    ("  CONTROLS  : 스위치/노브/버튼 1행 = 1개. control_id = panel_id.control_code", False),
    ("  POSITIONS : position 1행 = 1개. position_id = control_id.position_code  ← 절차/행위 문서는 이 ID를 참조", False),
    ("  GRID_<패널> : 패널 물리 배치 격자 (사람 검수용, LLM 조회 대상 아님). 셀에 control_id 병기", False),
    ("", False),
    ("2. 식별자 규칙", True),
    ("  panel_id      : 짧은 영문 대문자 (AUD1, IFF, ANTSEL, LG, ELEC, AVPWR …)", False),
    ("  control_code  : TO 명칭 기반 영문 대문자+밑줄 (COMM1_PWR, COMM1_MODE, M4_REPLY). 볼륨 전용 노브는 _VOL, 전원 겸용은 _PWR", False),
    ("  position_code : 패널 각인을 대문자+밑줄로 정규화 (OFF, SQL, GD, INC_ON, NONE_OFF, M3_MS, A_B, PUSH, NA). 숫자 position은 숫자 그대로", False),
    ("  예) AUD1.COMM1_MODE.GD = AUDIO 1 패널 · COMM 1 Mode Knob · GD 위치", False),
    ("  grid_row / grid_col : 패널 그리드 좌표 (1,1 = 좌상단). TO 그림 항목 번호는 to_item_no", False),
    ("  zone / access_rank : 패널·컨트롤·position 행 모두에 복사 — 어느 표에서든 접근 우선순위 조회 가능", False),
    ("", False),
    ("3. 표기 규칙", True),
    ("  볼륨 전용 노브(OFF 각인 없음) : NONE (OFF) / INC (ON) 2행 · 전원 겸용 노브 : OFF / INC (ON) 2행", False),
    ("  순간 푸시 버튼 : PUSH 1행 · 표시기/커넥터 : — (NA) 1행 · 숫자 썸휠 : 1행으로 축약 (note에 명시)", False),
    ("  TO 본문에 설명이 없는 항목은 패널 각인 기준으로 기재하고 note에 'TO 본문 설명 없음' 표시", False),
    ("  dcsbios_id / dcsbios_value 열은 패널 배치 완료 후 추가 예정", False),
    ("", False),
    ("  표시등(light) : OFF / ON 2행, function = 점등 의미. 연속 조정 노브/휠 : ADJ 1행", False),
    ("  N/I = DCS 미구현(Not Implemented), 비기능 = 해당 블록에서 기능 없음", False),
    ("", False),
    ("4. 출처", True),
    ("  DCS F-16C Early Access Guide EN, Cockpit Overview p.43-79 (주 출처)", False),
    ("  TO 1F-16CM/AM-1 BMS (CHANGE 4.37.4) p.148-160 — IFF/AUDIO/ANT SEL 최초 작성 후 DCS 가이드와 대조·수정. UHF BACKUP은 DCS 가이드 p.260-262 정의로 교체", False),
    ("  HUD/EHSI는 가이드의 별도 장 — 추후 추가. 방어 계통(p.693-708), HOTAS(p.83-89), UHF BACKUP·UHF/VHF DED(p.253-262) 추가 완료", False),
    ("", False),
    ("5. ICP / DED 3층 구조 (DCS EA Guide p.98-121)", True),
    ("  1층 물리 패널 : CONTROLS/POSITIONS의 ICP 패널. 키패드는 ICP.KEY_1~KEY_0 (각인은 control_name에만), 의미는 2층에서 결정", False),
    ("  2층 DED 그래프 : DED_PAGES(페이지 노드) · DED_TRANSITIONS(from_page + input → to_page) · DED_FIELDS(페이지별 필드, edit_method)", False),
    ("     edit_method = KEYPAD_ENTR(키패드 입력 후 ENTR) | MSEL_TOGGLE(0/M-SEL로 활성 토글) | KEY_TOGGLE(키패드 1-9 아무 키로 토글) | INCDEC_CYCLE(로커로 순환) | INCDEC_SPECIAL | SEQ_TOGGLE | DISPLAY(표시 전용)", False),
    ("     asterisk_default = 페이지 진입 시 별표(편집 커서)가 놓이는 필드. DCS UP/DN으로 field_no 순서대로 이동", False),
    ("  3층 조작 매크로 : ICP_ACTIONS — 절차 문서가 참조할 action_id, 전제 페이지, 입력 열(control_id 또는 position_id, > 구분), 검증(DED 텍스트)", False),
    ("  PANEL_ACTIONS : ICP/MFD가 아닌 일반 패널·HOTAS 조작 매크로 (CMDS 세팅, RWR 전원, ECM 모드, HOTAS 등) — 같은 열 구성", False),
    ("  HOTAS_FUNCTIONS : HOTAS 스위치는 position(방향)만 POSITIONS에 두고, 실제 기능은 (control, position, press, context) 조합으로 이 표에 기재", False),
    ("     press = SHORT(<0.5 s) | LONG(>0.5 s, 일부 1 s) | HOLD | ANY · context_type = MASTER_MODE | SOI | CMDS_MODE | STATE | ANY", False),
    ("     검증은 DCS-BIOS가 내보내는 DED 5행 텍스트를 읽어 수행 (스위치 상태가 아닌 화면 텍스트 기준)", False),
    ("", False),
    ("6. MFD 3층 구조 (DCS EA Guide p.122-128) — ICP와 같은 틀", True),
    ("  1층 물리 패널 : MFDL / MFDR (OSB_1~20 + GAIN/SYM/BRT/CON 로커). OSB 번호 = 상단 1→5, 우측 6→10, 하단 11→15(우→좌), 좌측 16→20(하→상)", False),
    ("  2층 포맷 그래프 : MFD_FORMATS(포맷 노드, master_menu_osb) · MFD_OSB_MAP(포맷별 OSB 의미, action_type) · MFD_TRANSITIONS", False),
    ("     OSB 11~15는 모든 포맷 공통(COMMON): 11 DCLT, 12/13/14 Format Select, 15 SWAP. Format Select 할당은 마스터모드 7종별로 독립", False),
    ("  3층 조작 매크로 : MFD_ACTIONS (ICP_ACTIONS와 같은 열 구성). 검증은 DCS-BIOS MFD OSB 라벨 텍스트", False),
    ("  OSB 맵 채운 포맷: MASTER_MENU, COMMON, DTE_P1/P2. FCR/TGP/WPN/SMS/HSD/HAD는 해당 장 읽을 때 보완", False),
    ("  현재 fields_filled=N 페이지(T-ILS/STPT/DEST/INS 등)는 해당 장을 읽을 때 DED_FIELDS 보완. CMDS 5페이지는 p.700-702로 보완 완료", False),
]
for i, (t, b) in enumerate(lines, start=1):
    ws0.cell(i, 1, t).font = f(b, size=12 if i == 1 else 10)
ws0.column_dimensions["A"].width = 120

write_table(wb, "ZONES", ["zone_id", "access_rank", "zone_name", "description"], [list(z) for z in ZONES], [10, 12, 48, 60])
write_table(wb, "PANELS", ["panel_id", "seq", "zone", "access_rank", "zone_seq", "panel_name", "location", "grid_rows", "grid_cols", "source", "note"], P,
            [12, 6, 8, 12, 9, 50, 30, 10, 10, 60, 60])
write_table(wb, "CONTROLS", ["control_id", "panel_id", "zone", "access_rank", "control_code", "to_item_no", "control_name", "control_type",
                             "grid_row", "grid_col", "grid_ref", "n_positions", "push_function", "note"], C,
            [24, 10, 8, 12, 16, 10, 40, 40, 9, 9, 9, 11, 50, 50])
write_table(wb, "POSITIONS", ["position_id", "control_id", "panel_id", "zone", "access_rank", "control_code", "position_code", "position_index",
                              "position_label", "function"], X,
            [32, 24, 10, 8, 12, 16, 14, 13, 16, 80])

write_table(wb, "DED_PAGES", ["page_id", "page_name", "group", "access_from", "access_input", "implemented", "fields_filled", "source_page", "note"], DED_PAGES,
            [12, 36, 10, 12, 14, 12, 12, 11, 50])
write_table(wb, "DED_TRANSITIONS", ["transition_id", "from_page", "input", "to_page", "note"], DED_TRANSITIONS,
            [13, 12, 16, 12, 60])
write_table(wb, "DED_FIELDS", ["field_id", "page_id", "field_no", "label", "editable", "edit_method", "value_format", "asterisk_default", "description", "dcs_note"], DED_FIELDS,
            [24, 11, 9, 16, 9, 16, 26, 15, 80, 22])
write_table(wb, "ICP_ACTIONS", ["action_id", "goal", "precondition_page", "input_sequence", "verify", "source_page", "note"], ICP_ACTIONS,
            [24, 40, 16, 70, 50, 11, 50])
write_table(wb, "HOTAS_FUNCTIONS", ["function_id", "control_id", "position", "press", "context_type", "context", "function", "source_page"], HOTAS_FUNCTIONS,
            [40, 24, 14, 8, 14, 20, 60, 11])
write_table(wb, "PANEL_ACTIONS", ["action_id", "goal", "precondition", "input_sequence", "verify", "source_page", "note"], PANEL_ACTIONS,
            [24, 40, 30, 70, 50, 11, 50])
write_table(wb, "MFD_FORMATS", ["format_id", "format_name", "master_menu_osb", "implemented", "osb_map_filled", "source_page", "description"], MFD_FORMATS,
            [14, 36, 16, 12, 14, 11, 70])
write_table(wb, "MFD_OSB_MAP", ["osb_map_id", "format_id", "osb_no", "label", "action_type", "function", "dcs_note"], MFD_OSB_MAP,
            [22, 14, 8, 18, 16, 90, 22])
write_table(wb, "MFD_TRANSITIONS", ["transition_id", "from_format", "input", "to_format", "note"], MFD_TRANSITIONS,
            [13, 14, 40, 28, 70])
write_table(wb, "MFD_ACTIONS", ["action_id", "goal", "precondition", "input_sequence", "verify", "source_page", "note"], MFD_ACTIONS,
            [24, 40, 24, 60, 50, 11, 50])

for p in PANELS:
    ws = write_panel(wb, p); ws.title = "GRID_" + p["code"]

out = "F-16_Cockpit_Panels_Duality.xlsx"
wb.save(out)
print("saved", out, "| panels", len(P), "controls", len(C), "positions", len(X))
